#!/usr/bin/env python3
"""
patch_eft.py — called by GitHub Actions with env vars set.

ENV VARS (all required):
  WORKBOOK       path to the xlsx file
  MATCH_JSON     JSON array of match objects:
                   [{"date":"06/05/2026","ref":"GERALDINE EVANS",
                     "amount":200,"row":123,"name":"Jane Smith"}]
  MONTH_EFT_COL  integer column number for the current month's EFT column
"""
import json, os, sys
import openpyxl

workbook_path = os.environ['WORKBOOK']
match_json    = os.environ['MATCH_JSON']
eft_col       = int(os.environ['MONTH_EFT_COL'])

matches = json.loads(match_json)
if not matches:
    print("No matches provided — nothing to do.")
    sys.exit(0)

wb  = openpyxl.load_workbook(workbook_path)
ws  = wb['Comprehensive list']
log = wb['Matching Audit Log']
nxt = log.max_row + 1

for m in matches:
    row    = m['row']
    amount = m['amount']
    ref    = m['ref'].strip()

    # 1. Write EFT amount into the current month's EFT column
    existing = ws.cell(row, eft_col).value or 0
    ws.cell(row, eft_col).value = existing + amount

    # 2. Permanently add the bank reference to column C (D/O Ref)
    #    so it auto-matches in all future months — never clear existing refs
    current_refs = ws.cell(row, 3).value or ''
    existing_tokens = [t.strip().upper() for t in current_refs.split(',') if t.strip()]
    if ref.upper() not in existing_tokens:
        new_refs = (current_refs.strip().rstrip(',') + ', ' + ref).strip(', ')
        ws.cell(row, 3).value = new_refs
        print(f"  D/O Ref updated: '{current_refs}' → '{new_refs}'")
    else:
        print(f"  D/O Ref already contains '{ref}' — skipping duplicate")

    # 3. Append to Matching Audit Log
    log.cell(nxt, 1).value = m['date']
    log.cell(nxt, 2).value = ref
    log.cell(nxt, 3).value = amount
    log.cell(nxt, 4).value = 'Membership (EFT)'
    log.cell(nxt, 5).value = 'MANUAL'
    log.cell(nxt, 6).value = 'HIGH'
    log.cell(nxt, 7).value = m['name']
    log.cell(nxt, 8).value = row
    log.cell(nxt, 9).value = 'Matched via dashboard'
    nxt += 1

    print(f"Patched: row {row} ({m['name']}) += R{amount}  [{ref}]")

wb.save(workbook_path)
print(f"Saved: {workbook_path}")
